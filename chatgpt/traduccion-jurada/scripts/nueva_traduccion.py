#!/usr/bin/env python3
"""
Genera el esqueleto de una traduccion jurada con el formato de la casa ya cargado.

    python3 nueva_traduccion.py --idioma en --docs "Birth Certificate" "Apostille" \
        --fecha 2026-08-25 --salida "20260825 - Traduccion Certificado Nacimiento (EN).docx"

Nunca partir de una traduccion anterior. Siempre de aqui.
"""
import argparse
import datetime as dt

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

# --- Formato de la casa -------------------------------------------------------
FUENTE = "Arial"
TAMANO = Pt(10)
INTERLINEADO = 1.5
MARGEN_IZQ = Inches(1.18)
MARGEN_DER = Inches(1.32)
MARGEN_SUP = Inches(1.0)
MARGEN_INF = Inches(2.25)
TAB_GUIONES = Inches(6.0)

MESES_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
            "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

CREDENCIALES_ES = (
    "Yo, {NOMBRE}, Traductora Jurada; debidamente autorizada en la República de Guatemala "
    "según Resolución número 346-DDEGN-2018 emitida por la Dirección de Educación "
    "Departamental Norte del Ministerio de Educación, de acuerdo con el Acuerdo Ministerial "
    "número 3850-2017 de fecha nueve de julio de dos mil dieciocho y con registro número "
    "004 – 2018, para traducir del idioma español a inglés y del inglés al español. Teniendo "
    "las traducciones fe pública, valor legal; declaro bajo juramento tener a la vista para "
    "traducir {ALCANCE} {ARTICULO}: "
)
CREDENCIALES_EN = (
    "I, {NOMBRE}, sworn translator; duly authorized in the Republic of Guatemala according "
    "to Resolution number 346-DDEGN-2018 issued by the Departmental Directorate of Education "
    "of the North of the Ministry of Education, in accordance with Ministerial Agreement "
    "Number 3850-2017 dated July 9, 2018 and with the registration number 004 - 2018, to "
    "translate with public faith and legal value from Spanish to English and from English "
    "into Spanish. I certify under oath that I have had at sight {ARTICULO}: "
)
NOMBRE = "Sofía Estela Rodríguez Farfán"
ROMANOS = ["i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x"]


def fecha_es(f):
    return f"{f.day} de {MESES_ES[f.month - 1]} de {f.year}"


def fecha_en(f):
    return f.strftime("%B %-d, %Y")


def configurar(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin, sec.right_margin = MARGEN_IZQ, MARGEN_DER
    sec.top_margin, sec.bottom_margin = MARGEN_SUP, MARGEN_INF

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = TAMANO
    # Arial tambien para caracteres no latinos, si no Word cae a otra fuente
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)

    pf = normal.paragraph_format
    pf.line_spacing = INTERLINEADO
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(0)

    pie(sec)
    return doc


def pie(sec):
    """Numero de pagina como campo PAGE en el pie."""
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, " PAGE "), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def parrafo(doc, partes=None, guiones=True):
    """Parrafo con tabulador derecho de guiones a 6.0. partes = [(texto, negrita), ...]"""
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(
        TAB_GUIONES, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DASHES
    )
    for texto, negrita in (partes or []):
        r = p.add_run(texto)
        r.bold = negrita
    if guiones:
        p.add_run("\t")
    return p


def enumeracion(docs):
    return "; ".join(f"{ROMANOS[i]}) {d}" for i, d in enumerate(docs))


def apertura(doc, idioma, docs, alcance_parcial):
    multiples = len(docs) > 1
    if idioma == "es":
        inicio = "Yo, "
        articulo = "los siguientes documentos" if multiples else "el siguiente documento"
        alcance = ("únicamente en su parte conducente" if alcance_parcial
                   else "de forma literal en su parte conducente")
        cuerpo = CREDENCIALES_ES.format(NOMBRE="", ALCANCE=alcance, ARTICULO=articulo)
        cola = (" los cuales se leen de la siguiente forma:" if multiples
                else " el cual se lee de la siguiente forma:")
    else:
        inicio = "I, "
        articulo = "the following documents" if multiples else "the following document"
        cuerpo = CREDENCIALES_EN.format(NOMBRE="", ARTICULO=articulo)
        cola = (", which I translate in their main part as follows:" if multiples
                else ", which I translate in its main part as follows:")

    # CREDENCIALES_* trae "Yo, {NOMBRE}, ..." / "I, {NOMBRE}, ..."; el nombre se
    # inserta aparte para que vaya en negrita.
    resto = cuerpo[len(inicio):].lstrip(", ")
    parrafo(doc, [
        (inicio, False),
        (NOMBRE, True),
        (", " + resto, False),
        (enumeracion(docs), True),
        (cola, False),
    ])


def cierre(doc, idioma, hojas, fecha):
    parrafo(doc, guiones=False)
    if idioma == "es":
        hoja = "hoja" if hojas == 1 else "hojas"
        impresa = "impresa" if hojas == 1 else "impresas"
        partes = [
            ("EN FE DE LO CUAL", True),
            (", sin asumir responsabilidad alguna por el contenido del documento traducido, "
             "a solicitud de la parte interesada, para los fines legales que le sean "
             "convenientes y bajo juramento, extiendo, sello y firmo esta traducción jurada "
             f"contenida en {hojas} {hoja} de papel bond, {impresa} únicamente en el lado "
             f"anverso. Guatemala, {fecha_es(fecha)}.", False),
        ]
    else:
        pagina = "bond page" if hojas == 1 else "bond pages"
        partes = [
            ("IN WITNESS WHEREOF", True),
            (", at the request of the interested party and for the uses that this party may "
             "deem convenient, I issue, sign, and seal the present translation in "
             f"{hojas} {pagina}, undertaking no responsibilities for the contents of the "
             f"original document. Issued in the City of Guatemala on {fecha_en(fecha)}.",
             False),
        ]
    parrafo(doc, partes)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--idioma", choices=["es", "en"], required=True,
                    help="idioma DESTINO de la traduccion")
    ap.add_argument("--docs", nargs="+", required=True,
                    help="nombre real de cada documento del lote, en orden")
    ap.add_argument("--hojas", type=int, default=0,
                    help="hojas de papel bond; se llena al final, 0 deja marcador")
    ap.add_argument("--fecha", default=None, help="AAAA-MM-DD (por defecto, hoy)")
    ap.add_argument("--parcial", action="store_true",
                    help="traduccion de parte conducente unicamente")
    ap.add_argument("--salida", required=True)
    a = ap.parse_args()

    fecha = dt.date.fromisoformat(a.fecha) if a.fecha else dt.date.today()
    doc = configurar(Document())
    apertura(doc, a.idioma, a.docs, a.parcial)

    for d in a.docs:
        parrafo(doc, [(f"[{d}]", True)])
        parrafo(doc, [("<< cuerpo del documento >>", False)])

    cierre(doc, a.idioma, a.hojas if a.hojas else "N", fecha)
    doc.save(a.salida)
    print(f"OK  {a.salida}")
    print("Recorda: contar las hojas reales y escribir ese numero en el cierre,")
    print("luego correr qa.py antes de entregar.")


if __name__ == "__main__":
    main()
